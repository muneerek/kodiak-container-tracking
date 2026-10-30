from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
import re
root=Path('docs/fsd')
def clean(s):
 s=re.sub(r'\[([^]]+)\]\([^)]+\)',r'\1',s)
 return s.replace('**','').replace('`','')
for src in sorted(root.glob('0*-FSD.md')):
 d=Document(); sec=d.sections[0];sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.65)
 for n in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
  st=d.styles[n];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
 d.styles['Normal'].font.size=Pt(10)
 d.styles['Normal'].paragraph_format.space_after=Pt(6)
 d.styles['Normal'].paragraph_format.line_spacing=1.08
 for n,size in [('Title',24),('Heading 1',15),('Heading 2',12)]: d.styles[n].font.size=Pt(size)
 lines=src.read_text(encoding='utf-8').splitlines();i=0
 while i<len(lines):
  s=lines[i]
  if not s.strip():i+=1;continue
  if s.startswith('```'):
   block=[];i+=1
   while i<len(lines) and not lines[i].startswith('```'):block.append(lines[i]);i+=1
   p=d.add_paragraph();p.paragraph_format.keep_together=True
   r=p.add_run('\n'.join(block));r.font.name='Consolas';r.font.size=Pt(8)
  elif s.startswith('|'):
   rows=[]
   while i<len(lines) and lines[i].startswith('|'):
    cells=[clean(c.strip()) for c in lines[i].strip().strip('|').split('|')]
    if not all(re.fullmatch(r'[-: ]+',c) for c in cells):rows.append(cells)
    i+=1
   i-=1
   t=d.add_table(rows=0,cols=len(rows[0]));t.autofit=False
   weights={2:[1.2,5.8],3:[1.65,2.15,3.2],4:[1.35,1.35,.9,3.4]}[len(rows[0])]
   for col,w in zip(t.columns,weights):col.width=Inches(w)
   for ri,row in enumerate(rows):
    cells=t.add_row().cells
    for cell,w,value in zip(cells,weights,row):
     cell.width=Inches(w);cell.text=value
     pr=cell._tc.get_or_add_tcPr(); shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCE6F1' if ri==0 else ('F5F7FA' if ri%2==0 else 'FFFFFF'));pr.append(shade)
     for p in cell.paragraphs:
      p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4)
      for r in p.runs:r.font.size=Pt(9);r.bold=ri==0
    trpr=t.rows[-1]._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
    if ri==0:trpr.append(OxmlElement('w:tblHeader'))
   borders=OxmlElement('w:tblBorders')
   for edge in ['top','left','bottom','right','insideH','insideV']:
    el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
   t._tbl.tblPr.append(borders)
   d.add_paragraph().paragraph_format.space_after=Pt(0)
  elif s.startswith('#'):
   level=len(s)-len(s.lstrip('#'));title=clean(s.lstrip('# ').strip());title=re.sub(r'[^\w\s]',' ',title);title=' '.join(title.split())
   d.add_paragraph(title,'Title' if level==1 else 'Heading '+str(min(level-1,3)))
  else:
   style='Normal'
   if s.startswith('- '):s=s[2:];style='List Bullet'
   elif re.match(r'^\d+\. ',s):s=re.sub(r'^\d+\. ','',s);style='List Number'
   d.add_paragraph(clean(s),style)
  i+=1
 p=sec.footer.paragraphs[0];p.alignment=2;p.add_run('Kodiak FSD  |  ')
 fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');p._p.append(fld)
 d.core_properties.author='Eknixis';d.core_properties.title=clean(lines[0].lstrip('# '))
 out=src.with_suffix('.docx');d.save(out);print(out)
